import json
import sqlite3
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Brak biblioteki openpyxl. Zainstaluj: pip install openpyxl")

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
DATA_DIR = PROJECT / "excel_do_wypelniania"
DB_PATH = HERE / "uczelnia.db"
SCHEMA_PATH = HERE / "schema.sql"
SEARCH_DATA_JS = PROJECT / "frontend" / "search-data.js"

BUILDING = {
    "name": "Wydział Inżynierii Mechanicznej i Mechatroniki (hala technologiczna „stara”/dydaktyczny)",
    "address": "ul. Sikorskiego 38, 70-323 Szczecin",
    "managers": "Małgorzata Jędrzejczak, Ewelina Kandyba",
    "transport": "Przystanki tramwajowe „Sikorskiego” – 170 m, „Narutowicza” – 300 m",
    "lat": 53.42362,
    "lon": 14.53483,
    "entrances_count": 6,
    "entrances_info": "Wejście boczne do Centrum Mechatroniki: schody z obustronnymi poręczami. "
                      "Pozostałe wejścia z poziomu terenu.",
    "automatic_doors": 0,
    "stairs_handrails": 1,
    "elevator": 0,
    "stair_lift": 0,
    "ramp": 0,
    "accessible_toilet": 0,
    "disabled_parking": "2 miejsca z boku budynku przy Wydziale Elektrycznym, "
                        "1 miejsce od strony budynku „nowa” hala technologiczna",
    "assistance_dog": 1,
    "sign_language": "Brak tłumacza języka migowego (także on-line). "
                     "Zapotrzebowanie zgłosić z 3-dniowym wyprzedzeniem.",
    "notes": None,
}

EXCEL_TABLES = {
    "rooms": ("sale.xlsx", ["id", "number", "wing", "floor", "capacity",
                            "owner", "room_type", "equipment", "notes"]),
    "lecturers": ("wykladowcy.xlsx", ["id", "name", "room_id", "notes"]),
}


def create_db(conn):
    conn.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
    cols = ["id"] + list(BUILDING)
    conn.execute(
        f"INSERT OR REPLACE INTO building ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
        [1] + list(BUILDING.values()),
    )
    print(f"Baza: {DB_PATH} (budynek wpisany)")


def cell_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def read_excel(path, columns):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = [cell_text(h) for h in next(rows_iter, [])]
    missing = [c for c in columns if c not in header]
    if missing:
        sys.exit(f"{path.name}: brak kolumn {missing} w pierwszym wierszu (jest: {header})")
    pos = {c: header.index(c) for c in columns}

    rows = []
    for r in rows_iter:
        row = [cell_text(r[pos[c]]) if pos[c] < len(r) else None for c in columns]
        if row[0]:
            rows.append(row)
    wb.close()
    return rows


def import_excel(conn, table):
    filename, columns = EXCEL_TABLES[table]
    path = DATA_DIR / filename
    if not path.exists():
        print(f"{table}: brak pliku {filename} - pomijam")
        return
    rows = read_excel(path, columns)
    ids = [r[0] for r in rows]
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    if dupes:
        sys.exit(f"{filename}: powtórzone id: {dupes} - każde id może wystąpić tylko raz")
    conn.execute(f"DELETE FROM {table}")
    conn.executemany(
        f"INSERT INTO {table} ({','.join(columns)}) VALUES ({','.join('?' * len(columns))})",
        rows,
    )
    print(f"{table}: {len(rows)} (z {filename})")


def floor_label(floor):
    try:
        f = int(floor)
    except (TypeError, ValueError):
        return None
    return "parter" if f == 0 else f"piętro {f}"


def generate_search_data(conn):
    conn.row_factory = sqlite3.Row
    items = []
    for r in conn.execute("SELECT * FROM rooms ORDER BY number"):
        sub = " · ".join(x for x in [
            (r["room_type"] or "").capitalize() or None,
            floor_label(r["floor"]),
            f"skrzydło {r['wing']}" if r["wing"] else None,
        ] if x)
        items.append({"type": "sala", "id": r["id"], "label": f"Sala {r['number']}",
                      "sub": sub, "keywords": r["room_type"] or ""})
    for r in conn.execute("SELECT * FROM lecturers ORDER BY rowid"):
        items.append({"type": "wykladowca", "id": r["id"], "label": r["name"],
                      "sub": f"Gabinet {r['room_id']}" if r["room_id"] else "",
                      "room": r["room_id"]})

    lines = ",\n".join("  " + json.dumps(i, ensure_ascii=False) for i in items)
    SEARCH_DATA_JS.write_text(
        f"window.SEARCH_DATA = [\n{lines}\n];\n",
        encoding="utf-8",
    )
    print(f"Wyszukiwarka: {len(items)} pozycji -> {SEARCH_DATA_JS.relative_to(PROJECT)}")


if __name__ == "__main__":
    with sqlite3.connect(DB_PATH) as conn:
        create_db(conn)
        for table in EXCEL_TABLES:
            import_excel(conn, table)
        generate_search_data(conn)
